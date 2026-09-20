from fastapi.testclient import TestClient
from io import BytesIO
import json

from docx import Document

from signalens import main
from signalens.triage import Decision


def _context():
    return {"research_areas": ["Semiconductor"], "watchlist": ["NVDA"], "investment_horizon": "6_24m"}


def test_json_flow_and_no_fourth_business_label(monkeypatch):
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-only")

    async def mock_jev(text, context, scope, key):
        assert text == "A supplier reports a verifiable bandwidth change."
        assert scope == "pasted_text"
        assert key == "test-only"
        return Decision("RESEARCH", 0.81, {"RESEARCH": 0.88, "KEEP": 0.12, "SKIP": 0.0}, "jev-1.13.0", 300)

    monkeypatch.setattr(main, "evaluate_with_jev", mock_jev)
    client = TestClient(main.app)
    response = client.post("/api/evaluate", json={
        "information_text": "A supplier reports a verifiable bandwidth change.",
        "research_context": _context(),
    })
    assert response.status_code == 200
    data = response.json()
    assert data["decision"] == "RESEARCH"
    assert data["confidence_score"] == 0.81
    assert data["confidence_kind"] == "jev_distribution_confidence"
    assert "information_text" not in data
    assert response.headers["cache-control"] == "no-store"


def test_invalid_context_never_returns_skip(monkeypatch):
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-only")
    client = TestClient(main.app)
    response = client.post("/api/evaluate", json={
        "information_text": "Some research information.",
        "research_context": {"research_areas": [], "watchlist": [], "investment_horizon": "6_24m"},
    })
    assert response.status_code == 400
    assert "decision" not in response.json()


def test_website_origin_cannot_call_local_api(monkeypatch):
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-only")
    client = TestClient(main.app)
    response = client.post("/api/evaluate", headers={"Origin": "https://random.example"}, json={
        "information_text": "Some research information.", "research_context": _context(),
    })
    assert response.status_code == 403
    assert "decision" not in response.json()


def test_missing_jev_key_is_explicit(monkeypatch):
    monkeypatch.delenv("TYPESAFE_API_KEY", raising=False)
    client = TestClient(main.app)
    response = client.post("/api/evaluate", json={
        "information_text": "Some research information.", "research_context": _context(),
    })
    assert response.status_code == 503
    assert response.json()["error_code"] == "EVALUATION_UNAVAILABLE"
    assert "decision" not in response.json()


def test_word_upload_uses_whole_document_scope(monkeypatch):
    monkeypatch.setenv("TYPESAFE_API_KEY", "test-only")
    document = Document()
    document.add_heading("New evidence", level=1)
    document.add_paragraph("A supplier reports higher accelerator interconnect demand.")
    stream = BytesIO()
    document.save(stream)

    async def mock_jev(text, context, scope, key):
        assert "interconnect demand" in text
        assert scope == "whole_document_text"
        return Decision("KEEP", 0.71, {"RESEARCH": 0.2, "KEEP": 0.7, "SKIP": 0.1}, "jev-1.13.0", 410)

    monkeypatch.setattr(main, "evaluate_with_jev", mock_jev)
    client = TestClient(main.app)
    response = client.post(
        "/api/evaluate",
        data={"research_context": json.dumps(_context())},
        files={"file": ("research.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers={"Origin": "chrome-extension://abcdefghijklmnopabcdefghijklmnop"},
    )
    assert response.status_code == 200
    result = response.json()
    assert result["decision"] == "KEEP"
    assert result["input_format"] == "docx"
    assert result["evaluated_scope"] == "whole_document_text"
    assert "file" not in result
