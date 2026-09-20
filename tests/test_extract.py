from io import BytesIO

import pymupdf
import pytest
from docx import Document

from signalens.extract import ExtractionError, extract_file, extract_pasted_text


def _pdf_bytes() -> bytes:
    pdf = pymupdf.open()
    page1 = pdf.new_page()
    page1.insert_text((50, 80), "Background information about the market.")
    page2 = pdf.new_page()
    page2.insert_text((50, 80), "A specific supplier reports higher bandwidth demand.")
    data = pdf.tobytes()
    pdf.close()
    return data


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Industry context", level=1)
    doc.add_paragraph("Background information.")
    doc.add_heading("New evidence", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Bandwidth demand rises"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_pdf_defaults_to_entire_document_and_tracks_pages():
    result = extract_file("research.pdf", _pdf_bytes())
    assert result.evaluated_scope == "whole_document_text"
    assert result.pages_evaluated == [1, 2]
    assert "Background" in result.text
    assert "higher bandwidth demand" in result.text


def test_pdf_page_selection_is_explicit():
    result = extract_file("research.pdf", _pdf_bytes(), page_range="2")
    assert result.evaluated_scope == "selected_pages"
    assert result.pages_evaluated == [2]
    assert "Background" not in result.text


def test_docx_heading_and_table_survive_extraction():
    result = extract_file("research.docx", _docx_bytes())
    assert result.evaluated_scope == "whole_document_text"
    assert "Industry context" in result.text
    assert "Bandwidth demand rises" in result.text
    section = extract_file("research.docx", _docx_bytes(), section_index=1)
    assert section.evaluated_scope == "selected_section"
    assert "New evidence" in section.text
    assert "Background information" not in section.text


def test_never_evaluates_a_url_as_a_page_fetch():
    with pytest.raises(ExtractionError) as exc:
        extract_pasted_text("https://example.com/news")
    assert exc.value.code == "INPUT_INVALID"


def test_misnamed_binary_is_rejected():
    with pytest.raises(ExtractionError) as exc:
        extract_file("report.pdf", b"not a pdf")
    assert exc.value.code == "UNSUPPORTED_FILE"


def test_long_pdf_can_select_pages_without_silent_truncation():
    pdf = pymupdf.open()
    for number in range(101):
        page = pdf.new_page()
        page.insert_text((50, 80), f"Research item {number + 1}: bandwidth demand changed.")
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(ExtractionError) as exc:
        extract_file("long.pdf", data)
    assert exc.value.code == "CONTENT_TOO_LONG"
    assert exc.value.details["available_page_range"] == "1-101"
    selected = extract_file("long.pdf", data, page_range="101")
    assert selected.evaluated_scope == "selected_pages"
    assert selected.pages_evaluated == [101]
    assert "Research item 101" in selected.text
